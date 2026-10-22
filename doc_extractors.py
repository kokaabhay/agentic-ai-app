
import io
import csv
from pathlib import Path

from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook


SUPPORTED_EXTENSIONS = {
    ".pdf", ".docx", ".txt", ".md", ".csv", ".xlsx"
}


def detect_file_type(file_content: bytes, filename: str) -> str:
    extension = Path(filename).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type: {extension}. "
            "Supported formats: PDF, DOCX, TXT, MD, CSV, XLSX."
        )

    if extension == ".pdf":
        if not file_content.startswith(b"%PDF"):
            raise ValueError("Invalid PDF file.")

    elif extension in {".docx", ".xlsx"}:
        if not file_content.startswith(b"PK"):
            raise ValueError(f"Invalid {extension} file.")

    else:
        try:
            file_content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError("Text file must use UTF-8 encoding.") from exc

    return extension


def extract_pdf_text(file_content: bytes) -> str:
    reader = PdfReader(io.BytesIO(file_content))
    return "\n\n".join(
        page.extract_text() or "" for page in reader.pages
    ).strip()


def extract_docx_text(file_content: bytes) -> str:
    document = Document(io.BytesIO(file_content))
    parts = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    for table in document.tables:
        for row in table.rows:
            parts.append(
                " | ".join(cell.text.strip() for cell in row.cells)
            )

    return "\n".join(parts).strip()


def extract_text_file(file_content: bytes) -> str:
    return file_content.decode("utf-8-sig").strip()


def extract_csv_text(file_content: bytes) -> str:
    text = file_content.decode("utf-8-sig")
    rows = csv.reader(io.StringIO(text))

    return "\n".join(
        " | ".join(cell.strip() for cell in row)
        for row in rows
    ).strip()


def extract_xlsx_text(file_content: bytes) -> str:
    parts = []
    workbook = load_workbook(
        filename=io.BytesIO(file_content),
        read_only=True,
        data_only=True,
    )

    try:
        for worksheet in workbook.worksheets:
            parts.append(f"Sheet: {worksheet.title}")

            for row in worksheet.iter_rows(values_only=True):
                values = [
                    str(value).strip()
                    for value in row
                    if value is not None
                ]

                if values:
                    parts.append(" | ".join(values))
    finally:
        workbook.close()

    return "\n".join(parts).strip()


def extract_text(file_content: bytes, file_type: str) -> str:
    if file_type == ".pdf":
        text = extract_pdf_text(file_content)
    elif file_type == ".docx":
        text = extract_docx_text(file_content)
    elif file_type in {".txt", ".md"}:
        text = extract_text_file(file_content)
    elif file_type == ".csv":
        text = extract_csv_text(file_content)
    elif file_type == ".xlsx":
        text = extract_xlsx_text(file_content)
    else:
        raise ValueError(f"No extractor available for {file_type}.")

    if not text.strip():
        raise ValueError("No readable text or data found in the file.")

    return text


def extract_uploaded_file(filename: str, file_content: bytes) -> str:
    file_type = detect_file_type(file_content, filename)
    return extract_text(file_content, file_type)
